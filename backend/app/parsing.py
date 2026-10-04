import io
import re
import zipfile
from uuid import uuid4
from pypdf import PdfReader
from docx import Document

def parse(content: bytes, suffix: str) -> list[dict]:
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(content))
        if len(reader.pages) > 200:
            raise ValueError("Documents may contain at most 200 pages.")
        pages = [p.extract_text() or "" for p in reader.pages]
    elif suffix == ".docx":
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError("DOCX decompressed content exceeds 30 MB.")
        doc = Document(io.BytesIO(content))
        pages = ["\n".join([p.text for p in doc.paragraphs] + [" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows])]
    elif suffix == ".txt":
        pages = content.decode("utf-8-sig").split("\f")
    else:
        raise ValueError("Use a PDF, DOCX, or UTF-8 TXT document.")
    clauses = []
    for page, text in enumerate(pages, 1):
        parts = re.split(r"(?m)(?=^\s*\d+(?:\.\d+)*[.)]?\s+[A-Z])", text)
        for part in parts:
            part = part.strip()
            if not part:
                continue
            heading = part.splitlines()[0][:160]
            # Bound unusually long sections while preserving page and heading.
            for start in range(0, len(part), 6000):
                clauses.append({"id": str(uuid4()), "page": page, "heading": heading, "text": part[start:start+6000]})
    if not clauses:
        raise ValueError("No readable text found. Scanned PDFs need OCR before upload.")
    if sum(len(c["text"]) for c in clauses) > 1_000_000:
        raise ValueError("Extracted text exceeds the 1 MB limit.")
    return clauses

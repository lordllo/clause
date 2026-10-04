import io
from pathlib import Path
import pytest
from app.parsing import parse
from app.analysis import analyze, verify, retrieve
from app.schemas import DEFAULT_POLICY, Finding

FIXTURES = Path(__file__).resolve().parents[2] / "fixtures"

@pytest.mark.parametrize("name,status", [("vendor-risky.txt", "deviation"), ("vendor-compliant.txt", "compliant")])
def test_baseline(name, status, monkeypatch):
    monkeypatch.setenv("ANALYSIS_PROVIDER", "demo")
    clauses = parse((FIXTURES / name).read_bytes(), ".txt")
    result = analyze(DEFAULT_POLICY, clauses)
    assert len(result["findings"]) == 5
    assert all(f["status"] == status and f["citation_verified"] for f in result["findings"])

def test_forged_citation():
    c = {"id":"a", "text":"The laws of New York apply.", "page":2,"heading":"Law"}
    f = Finding(rule_id="governing_law",status="compliant",explanation="ok",action="none",clause_id="a",quote="The laws of California apply.")
    assert verify(f, [c])["status"] == "needs_review"
    assert not verify(f, [c])["citation_verified"]

def test_unknown_source_and_short_quote():
    f = Finding(rule_id="law",status="compliant",explanation="ok",action="none",clause_id="invented",quote="New York")
    assert not verify(f, [])["citation_verified"]

def test_missing_clause(monkeypatch):
    monkeypatch.setenv("ANALYSIS_PROVIDER", "demo")
    assert all(f["status"] == "needs_review" for f in analyze(DEFAULT_POLICY, [])['findings'])

def test_custom_policy_abstains(monkeypatch):
    monkeypatch.setenv("ANALYSIS_PROVIDER", "demo")
    policy = DEFAULT_POLICY.model_copy(deep=True)
    policy.rules[0].instruction = "Governing law must be California."
    clauses = parse((FIXTURES/'vendor-risky.txt').read_bytes(),'.txt')
    assert analyze(policy, clauses)['findings'][0]['status'] == 'needs_review'

def test_page_provenance():
    clauses = parse(b"1. Law\nNew York\f2. Security\n72 hours", ".txt")
    assert [c['page'] for c in clauses] == [1,2]

def test_docx():
    from docx import Document
    d=Document();d.add_paragraph('1. Governing Law');d.add_paragraph('New York applies.')
    stream=io.BytesIO();d.save(stream)
    assert 'New York' in parse(stream.getvalue(),'.docx')[0]['text']

def test_empty_rejected():
    with pytest.raises(ValueError): parse(b'', '.txt')
